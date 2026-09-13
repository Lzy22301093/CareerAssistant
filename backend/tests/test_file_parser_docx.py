"""Tests for FileParserTool DOCX extraction (tables / Chinese punctuation)."""

from __future__ import annotations

import io

import pytest

from app.tools.file_tools import FileParserTool, extract_docx_text


def _docx_with_table() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph("栗兆洋 - AI Agent 开发工程师")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "教育背景"
    table.cell(0, 1).text = "北京交通大学（211）· 软件工程"
    table.cell(1, 0).text = "工作内容"
    table.cell(1, 1).text = "采用 LangGraph，检索准确率提升 40%；覆盖数据采集 → 文档解析"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def test_extract_docx_from_table():
    from docx import Document

    doc = Document(io.BytesIO(_docx_with_table()))
    text = extract_docx_text(doc)
    assert "教育背景" in text
    assert "北京交通大学" in text
    assert "LangGraph" in text


async def test_parse_docx_bytes_from_table():
    tool = FileParserTool()
    import base64

    result = await tool.execute(
        file_content=base64.b64encode(_docx_with_table()).decode(),
        filename="resume.docx",
    )
    assert result.success, result.error
    text = result.data["text"]
    assert "教育背景" in text
    assert "LangGraph" in text
    # 中文标点与箭头不应被吞掉
    assert "，" in text or "·" in text
    assert "→" in text or "40%" in text


def test_clean_text_keeps_chinese_punct():
    tool = FileParserTool()
    cleaned = tool._clean_text("熟悉 Java、Spring；成绩优秀，能独立完成。路径：A → B · C")
    assert "、" in cleaned
    assert "；" in cleaned
    assert "，" in cleaned
    assert "：" in cleaned or ":" in cleaned
    assert "→" in cleaned
    assert "·" in cleaned


async def test_parse_docx_empty_raises():
    from docx import Document

    empty = Document()
    buf = io.BytesIO()
    empty.save(buf)
    tool = FileParserTool()
    import base64

    result = await tool.execute(
        file_content=base64.b64encode(buf.getvalue()).decode(),
        filename="empty.docx",
    )
    assert not result.success
    assert "未提取" in (result.error or "")
