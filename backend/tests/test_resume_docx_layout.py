"""Tests for resume original storage + DOCX layout preserve export."""

from __future__ import annotations

import io

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker

from app.models.database import Base
from app.models.orm import User
from app.models.schemas import ResumeDocumentCreate, ResumeVersionCreate
from app.services.resume_docx_layout import (
    DocxLayoutError,
    DocxLayoutService,
    ResumeSourceFileService,
    _norm_text,
)
from app.services.resume_library_service import ResumeLibraryService


@pytest.fixture
def db_session():
    engine = create_engine("sqlite:///:memory:", echo=False)
    Base.metadata.create_all(bind=engine)
    TestSession = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    session = TestSession()
    session.add(User(username="u1", email="u1@test.com", hashed_password="x"))
    session.commit()
    try:
        yield session
    finally:
        session.close()
        Base.metadata.drop_all(bind=engine)


def _make_min_docx(paragraphs: list[str]) -> bytes:
    from docx import Document

    doc = Document()
    for t in paragraphs:
        doc.add_paragraph(t)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_norm_text():
    assert _norm_text("  Hello，World！ ") == "hello,world!"
    assert _norm_text("（测试）") == "(测试)"


def test_save_and_get_source_file(db_session: Session, tmp_path, monkeypatch):
    monkeypatch.setattr("app.services.resume_docx_layout.settings.upload_dir", str(tmp_path))
    svc = ResumeSourceFileService()
    lib = ResumeLibraryService()
    doc = lib.create_document(db_session, 1, ResumeDocumentCreate(title="简历"))
    raw = _make_min_docx(["北京交通大学 软件工程", "项目经历"])
    row = svc.save_original(db_session, 1, doc.id, "me.docx", raw, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert row.id and row.sha256
    got = svc.get_for_document(db_session, 1, doc.id)
    assert got is not None and svc.is_docx(got)
    assert not svc.is_docx(None)


def test_layout_export_replaces_matched_section(db_session: Session, tmp_path, monkeypatch):
    monkeypatch.setattr("app.services.resume_docx_layout.settings.upload_dir", str(tmp_path))
    lib = ResumeLibraryService()
    src_svc = ResumeSourceFileService()
    doc = lib.create_document(db_session, 1, ResumeDocumentCreate(title="简历", source="upload"))
    original_line = "北京交通大学 软件工程 本科 2022-2026"
    raw = _make_min_docx(["姓名 张三", original_line, "技能 Python"])
    src_svc.save_original(db_session, 1, doc.id, "a.docx", raw)
    lib.add_version(
        db_session,
        1,
        doc.id,
        ResumeVersionCreate(
            content={
                "sections": [
                    {"title": "教育背景", "content": "北京交通大学 软件工程 本科 2022-2026（已改）"},
                    {"title": "技能", "content": "技能 Python FastAPI"},
                ],
                "raw_text": "",
            }
        ),
    )
    blob = DocxLayoutService().export_preserve(db_session, 1, doc.id)
    assert blob[:2] == b"PK"
    from docx import Document

    out = Document(io.BytesIO(blob))
    texts = "\n".join(p.text for p in out.paragraphs)
    assert "已改" in texts
    assert "FastAPI" in texts


def test_layout_export_no_source_raises(db_session: Session):
    lib = ResumeLibraryService()
    doc = lib.create_document(db_session, 1, ResumeDocumentCreate(title="无原件"))
    lib.add_version(db_session, 1, doc.id, ResumeVersionCreate(content={"sections": [{"title": "A", "content": "B"}], "raw_text": ""}))
    with pytest.raises(DocxLayoutError):
        DocxLayoutService().export_preserve(db_session, 1, doc.id)
