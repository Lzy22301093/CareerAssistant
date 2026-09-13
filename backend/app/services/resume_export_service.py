"""ResumeExportService — 简历生成区导出（Word/HTML/Markdown/JSON）。

- HTML：A4 打印友好模板，可嵌证件照（base64），浏览器可另存 PDF
- DOCX：python-docx，层级标题 + 段落 + 可选证件照
- MD：保留（结构化 `##` 输出）
- JSON：原样
正文统一去 Markdown 残留，避免 `###`/`- ` 直接进 Word/HTML。
"""

from __future__ import annotations

import base64
import html as _html
import io
import json
import re
from typing import Any

from docx.oxml.ns import qn
from docx.shared import Cm, Pt

VALID_FORMATS = {"docx", "html", "md", "markdown", "json"}

_MD_NOISE = re.compile(r"^\s{0,3}(#{1,6}\s+|\*\*|__|`{1,3})")
_MD_LIST = re.compile(r"^\s{0,3}[-*+]\s+")


class ResumeExportError(ValueError):
    """导出业务错误。"""


def strip_markdown(text: str) -> str:
    """去掉行首 Markdown 标记，保留语义文本。"""
    if not text:
        return ""
    lines: list[str] = []
    for line in str(text).splitlines():
        s = _MD_NOISE.sub("", line)
        s = _MD_LIST.sub("· ", s)
        s = s.replace("**", "").replace("__", "")
        lines.append(s.rstrip())
    return "\n".join(lines).strip()


def _normalized_sections(content: dict[str, Any]) -> list[dict[str, str]]:
    sections = content.get("sections") if isinstance(content, dict) else None
    if not isinstance(sections, list):
        return []
    out: list[dict[str, str]] = []
    for sec in sections:
        if not isinstance(sec, dict):
            continue
        title = strip_markdown(str(sec.get("title") or ""))
        body = strip_markdown(str(sec.get("content") or ""))
        if not title and not body:
            continue
        out.append({"title": title, "content": body})
    return out


# 导入拍平后的子块标题：导出时并回上一个父块，避免「项目描述/工作内容」与「项目经历」平级
_FLATTEN_CHILD_TITLES = {
    "项目描述",
    "工作内容",
    "工作职责",
    "岗位职责",
    "主要工作",
    "实习内容",
    "实习职责",
    "项目内容",
    "个人职责",
    "工作业绩",
    "项目成果",
}


def compose_sections_for_export(sections: list[dict[str, str]]) -> list[dict[str, str]]:
    """合并被标题行拍平的子区域（B），只影响导出组装，不改库内 section。"""
    out: list[dict[str, str]] = []
    for sec in sections:
        title = (sec.get("title") or "").strip()
        body = (sec.get("content") or "").strip()
        if title in _FLATTEN_CHILD_TITLES and out:
            parent = out[-1]
            piece = f"{title}：\n{body}" if body else f"{title}："
            parent["content"] = f"{parent['content']}\n{piece}".strip() if parent["content"] else piece
            continue
        out.append({"title": title, "content": body})
    return out


def prepare_export_sections(content: dict[str, Any]) -> list[dict[str, str]]:
    return compose_sections_for_export(_normalized_sections(content))


def _photo_data_url(photo_bytes: bytes | None, mime: str | None = None) -> str | None:
    if not photo_bytes:
        return None
    mime = (mime or "image/jpeg").split(";")[0].strip() or "image/jpeg"
    b64 = base64.b64encode(photo_bytes).decode("ascii")
    return f"data:{mime};base64,{b64}"


def _is_entry_header(line: str) -> bool:
    """经历块首行：含 | 分隔或「短标题 + 年份」。避免把正文短句误加粗。"""
    if line.startswith(("·", "•", "-")):
        return False
    if line.startswith(("职责", "成果", "主修课程", "工作内容", "项目描述", "求职意向")):
        return False
    if "|" in line:
        return True
    # 含年份时间段且较短 → 项目/实习标题行
    if re.search(r"20\d{2}", line) and len(line) <= 64:
        return True
    return False


# ---------- HTML（A4 打印模板） ----------

def build_html(
    content: dict[str, Any],
    title: str,
    photo_bytes: bytes | None = None,
    photo_mime: str | None = None,
) -> str:
    sections = prepare_export_sections(content)
    name = title or "个人简历"
    # 尝试从「基本信息」里取姓名作页眉
    for sec in sections:
        if sec["title"] == "基本信息":
            for line in sec["content"].splitlines():
                if line.startswith("姓名："):
                    name = line.split("：", 1)[-1].strip() or name
            break

    photo_url = _photo_data_url(photo_bytes, photo_mime)
    photo_html = (
        f'<img class="photo" src="{photo_url}" alt="证件照" />' if photo_url else ""
    )

    body_parts: list[str] = []
    for sec in sections:
        if sec["title"] == "基本信息":
            lines = [ln for ln in sec["content"].splitlines() if ln.strip()]
            contact = " · ".join(
                ln.split("：", 1)[-1].strip() if "：" in ln else ln.strip()
                for ln in lines
                if not ln.startswith("姓名：")
            )
            body_parts.append(
                f'<header class="hd"><div class="hd-main"><h1>{_html.escape(name)}</h1>'
                f'<p class="contact">{_html.escape(contact)}</p></div>{photo_html}</header>'
            )
            continue
        if sec["title"] == "求职意向":
            body_parts.append(
                f'<p class="obj">求职意向：{_html.escape(sec["content"])}</p>'
            )
            continue
        para_parts: list[str] = []
        for line in sec["content"].splitlines():
            if not line.strip():
                continue
            if line.startswith("·"):
                cls = "bullet"
            elif _is_entry_header(line):
                cls = "entry"
            else:
                cls = ""
            attr = f' class="{cls}"' if cls else ""
            para_parts.append(f"<p{attr}>{_html.escape(line)}</p>")
        paras = "".join(para_parts)
        body_parts.append(
            f'<section class="sec"><h2>{_html.escape(sec["title"])}</h2>{paras}</section>'
        )

    body = "\n".join(body_parts)
    return (
        '<!doctype html><html lang="zh-CN"><head><meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width, initial-scale=1">'
        f"<title>{_html.escape(title or name)}</title>"
        "<style>"
        "@page { size: A4; margin: 16mm 14mm; }"
        "* { box-sizing: border-box; }"
        "body { max-width: 210mm; margin: 0 auto; padding: 24px 28px; "
        "font-family: 'Microsoft YaHei','微软雅黑','PingFang SC','Noto Sans SC',sans-serif; "
        "color: #1f1c19; line-height: 1.55; background: #fff; font-size: 10.5pt; }"
        ".hd { display: flex; justify-content: space-between; align-items: flex-start; "
        "border-bottom: 2px solid #c15f3c; padding-bottom: 12px; margin-bottom: 16px; }"
        ".hd h1 { margin: 0; font-size: 22pt; letter-spacing: 0.04em; font-weight: 700; }"
        ".contact { margin: 6px 0 0; color: #444; font-size: 10pt; line-height: 1.45; }"
        ".photo { width: 92px; height: 124px; object-fit: cover; border: 1px solid #ddd; "
        "border-radius: 4px; margin-left: 16px; }"
        ".sec { margin: 10px 0 0; break-inside: avoid; }"
        ".sec h2 { margin: 0 0 4px; font-size: 11.5pt; color: #8b3a22; font-weight: 700; "
        "border-bottom: 1px solid #d9b5a4; padding-bottom: 3px; letter-spacing: 0.04em; }"
        ".sec p { margin: 1px 0; white-space: pre-wrap; font-size: 10pt; line-height: 1.45; }"
        ".sec p.bullet { padding-left: 8px; text-indent: -0.4em; }"
        ".sec p.entry { margin-top: 6px; font-weight: 700; font-size: 10.5pt; }"
        ".obj { margin: 0 0 10px; color: #8b3a22; font-weight: 700; font-size: 10pt; }"
        "@media print { body { padding: 0; max-width: none; } .photo { -webkit-print-color-adjust: exact; } }"
        "</style></head>"
        f"<body>{body}</body></html>"
    )


# ---------- DOCX ----------

# 统一字体：微软雅黑中英混排最稳（Windows Word 预装）；标题加粗同族，避免黑体/宋体混跳
FONT_BODY = "微软雅黑"
FONT_HEAD = "微软雅黑"
# 兼容旧常量
FONT_HEI = FONT_HEAD
FONT_SONG = FONT_BODY
_ACCENT_RGB = (0x1F, 0x1C, 0x19)
_ACCENT_TITLE_RGB = (0x8B, 0x3A, 0x22)  # 陶土深色标题
_GRAY_RGB = (0x2A, 0x2A, 0x2A)
_MUTED_RGB = (0x55, 0x55, 0x55)


def _apply_rfonts(rPr, name: str) -> None:
    """强制 latin/eastAsia/cs 同字体，并去掉主题字体覆盖。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    rFonts = rPr.get_or_add_rFonts()
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        key = qn(f"w:{attr}")
        if key in rFonts.attrib:
            del rFonts.attrib[key]
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rFonts.set(qn(f"w:{attr}"), name)
    # lang 保证中西文都走同一校对
    lang = rPr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rPr.append(lang)
    lang.set(qn("w:val"), "zh-CN")
    lang.set(qn("w:eastAsia"), "zh-CN")


def _set_run_font(run, *, name: str = FONT_BODY, size_pt: float = 10.5, bold: bool = False, color=None) -> None:
    run.bold = bold
    run.font.size = Pt(size_pt)
    if color is not None:
        run.font.color.rgb = color
    run.font.name = name
    _apply_rfonts(run._element.get_or_add_rPr(), name)


def _set_para_spacing(p, *, before: float = 0, after: float = 0, line: float = 1.15, left_cm: float = 0) -> None:
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if left_cm:
        pf.left_indent = Cm(left_cm)


def _set_para_bottom_border(p, *, color: str = "C15F3C", size: int = 8) -> None:
    """段落下边框（分区细线），比 ─ 字符更稳。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    pPr = p._element.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    bottom = pBdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        pBdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)


def _remove_table_borders(table) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else tbl.add_tblPr()
    borders = tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "nil")


def _parse_contact_fields(basic_lines: list[str]) -> list[tuple[str, str]]:
    """从基本信息行解析 (标签, 值)，避免长串 contact 在 Word 里从中间断行。"""
    pairs: list[tuple[str, str]] = []
    for ln in basic_lines:
        ln = ln.strip()
        if not ln or ln.startswith("姓名："):
            continue
        if ln.startswith("联系方式："):
            rest = ln.split("：", 1)[-1]
            for part in re.split(r"[·|｜]", rest):
                part = part.strip()
                if "：" in part:
                    k, v = part.split("：", 1)
                    pairs.append((k.strip(), v.strip()))
                elif part:
                    pairs.append(("", part))
            continue
        if "：" in ln:
            k, v = ln.split("：", 1)
            pairs.append((k.strip(), v.strip()))
        else:
            pairs.append(("", ln.strip()))
    return [(k, v) for k, v in pairs if v]


def _format_date_bit(bit: str) -> bool:
    return bool(re.search(r"\d{4}", bit)) and ("–" in bit or "-" in bit or "至今" in bit or "年" in bit)


def build_docx(
    content: dict[str, Any],
    title: str,
    photo_bytes: bytes | None = None,
    photo_mime: str | None = None,
) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    # 文档默认字体：中英统一，避免 Word 主题 Calibri 混入
    style = doc.styles["Normal"]
    style.font.name = FONT_BODY
    style.font.size = Pt(10.5)
    rPr = style.element.get_or_add_rPr()
    _apply_rfonts(rPr, FONT_BODY)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = 1.2

    for section in doc.sections:
        section.top_margin = Cm(1.4)
        section.bottom_margin = Cm(1.4)
        section.left_margin = Cm(1.6)
        section.right_margin = Cm(1.6)
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)

    sections = prepare_export_sections(content)
    name = title or "个人简历"
    basic_lines: list[str] = []
    objective = ""
    for sec in sections:
        if sec["title"] == "基本信息":
            basic_lines = [ln for ln in sec["content"].splitlines() if ln.strip()]
            for ln in basic_lines:
                if ln.startswith("姓名："):
                    name = ln.split("：", 1)[-1].strip() or name
        if sec["title"] == "求职意向":
            objective = sec["content"].strip()

    contact_pairs = _parse_contact_fields(basic_lines)
    ink = RGBColor(*_ACCENT_RGB)
    title_ink = RGBColor(*_ACCENT_TITLE_RGB)
    gray = RGBColor(*_GRAY_RGB)
    muted = RGBColor(*_MUTED_RGB)

    # 头区：左信息 / 右照片
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    try:
        table.columns[0].width = Cm(13.6)
        table.columns[1].width = Cm(3.4)
        table.cell(0, 0).width = Cm(13.6)
        table.cell(0, 1).width = Cm(3.4)
    except Exception:
        pass
    _remove_table_borders(table)

    left = table.cell(0, 0)
    # 清掉默认空段再写
    left_p = left.paragraphs[0]
    _set_para_spacing(left_p, after=2, line=1.0)
    _set_run_font(left_p.add_run(name), name=FONT_HEAD, size_pt=20, bold=True, color=ink)

    # 联系方式：每行 2 字段 + 制表位，避免「城市：北京」被拆行
    if contact_pairs:
        from docx.enum.text import WD_TAB_ALIGNMENT

        for i in range(0, len(contact_pairs), 2):
            chunk = contact_pairs[i : i + 2]
            cp = left.add_paragraph()
            _set_para_spacing(cp, after=1, line=1.2)
            try:
                cp.paragraph_format.tab_stops.add_tab_stop(Cm(6.8), WD_TAB_ALIGNMENT.LEFT)
            except Exception:
                pass
            for ci, (k, v) in enumerate(chunk):
                if ci:
                    _set_run_font(cp.add_run("\t"), name=FONT_BODY, size_pt=9, color=gray)
                if k:
                    _set_run_font(cp.add_run(f"{k}："), name=FONT_BODY, size_pt=9, color=muted)
                _set_run_font(cp.add_run(v), name=FONT_BODY, size_pt=9, color=gray)

    if objective:
        op = left.add_paragraph()
        _set_para_spacing(op, before=4, after=0, line=1.15)
        _set_run_font(op.add_run("求职意向："), name=FONT_HEAD, size_pt=10, bold=True, color=title_ink)
        _set_run_font(op.add_run(objective), name=FONT_BODY, size_pt=10, color=gray)

    right = table.cell(0, 1)
    if photo_bytes:
        try:
            rp = right.paragraphs[0]
            rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _set_para_spacing(rp, line=1.0)
            rrun = rp.add_run()
            rrun.add_picture(io.BytesIO(photo_bytes), width=Cm(2.6))
        except Exception:
            pass
    else:
        right.paragraphs[0].text = ""

    # 头区底线
    rule = doc.add_paragraph()
    _set_para_spacing(rule, before=2, after=6, line=1.0)
    _set_para_bottom_border(rule, color="C15F3C", size=12)

    for sec in sections:
        if sec["title"] in ("基本信息", "求职意向"):
            continue
        if not sec["title"] and not sec["content"]:
            continue

        lines = [ln.strip() for ln in sec["content"].splitlines() if ln.strip()]
        if not lines:
            continue

        if sec["title"]:
            h = doc.add_paragraph()
            _set_para_spacing(h, before=8, after=3, line=1.0)
            _set_run_font(h.add_run(sec["title"]), name=FONT_HEAD, size_pt=11.5, bold=True, color=title_ink)
            _set_para_bottom_border(h, color="D9B5A4", size=6)

        i = 0
        while i < len(lines):
            line = lines[i]
            # 经历块：标题行加粗 + 后续 bullet
            if _is_entry_header(line):
                p = doc.add_paragraph()
                _set_para_spacing(p, before=5, after=1, line=1.12)
                if "|" in line:
                    bits = [b.strip() for b in re.split(r"\s*\|\s*", line) if b.strip()]
                    for bi, bit in enumerate(bits):
                        if bi:
                            sep = p.add_run("  |  ")
                            _set_run_font(sep, name=FONT_BODY, size_pt=9, color=RGBColor(0x99, 0x99, 0x99))
                        if _format_date_bit(bit):
                            _set_run_font(p.add_run(bit), name=FONT_BODY, size_pt=9, color=muted)
                        else:
                            _set_run_font(p.add_run(bit), name=FONT_HEAD, size_pt=10.5, bold=True, color=ink)
                else:
                    _set_run_font(p.add_run(line), name=FONT_HEAD, size_pt=10.5, bold=True, color=ink)
                i += 1
                continue

            is_bullet = line.startswith("·") or line.startswith("•")
            text = line.lstrip("·• ").strip() if is_bullet else line
            p = doc.add_paragraph()
            if is_bullet:
                _set_para_spacing(p, after=1, line=1.2, left_cm=0.32)
                _set_run_font(p.add_run("· "), name=FONT_BODY, size_pt=10, color=title_ink)
                _set_run_font(p.add_run(text), name=FONT_BODY, size_pt=10, color=gray)
            else:
                # 「职责：xxx」「成果：xxx」标签加重
                if "：" in text[:6]:
                    lab, rest = text.split("：", 1)
                    _set_para_spacing(p, after=1, line=1.2)
                    _set_run_font(p.add_run(f"{lab}："), name=FONT_HEAD, size_pt=10, bold=True, color=gray)
                    _set_run_font(p.add_run(rest), name=FONT_BODY, size_pt=10, color=gray)
                else:
                    _set_para_spacing(p, after=1, line=1.2)
                    _set_run_font(p.add_run(text), name=FONT_BODY, size_pt=10, color=gray)
            i += 1

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ---------- MD / JSON ----------

def build_markdown(content: dict[str, Any], title: str) -> str:
    """Markdown 导出（保留）。正文先去噪音，再以 ## 标题输出。"""
    sections = prepare_export_sections(content)
    if sections:
        raw = "\n\n".join(f"## {s['title']}\n{s['content']}" for s in sections if s["title"])
    else:
        raw = strip_markdown(str(content.get("raw_text") or ""))
    return f"# {title or '简历'}\n\n{raw}".strip() + "\n"


def build_json(content: dict[str, Any]) -> str:
    return json.dumps(content, ensure_ascii=False, indent=2)


def build_export(
    content: dict[str, Any],
    title: str,
    format: str,
    photo_bytes: bytes | None = None,
    photo_mime: str | None = None,
) -> tuple[bytes, str, str]:
    fmt = (format or "docx").lower()
    if fmt == "docx":
        return (
            build_docx(content, title, photo_bytes, photo_mime),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "docx",
        )
    if fmt == "html":
        return (
            build_html(content, title, photo_bytes, photo_mime).encode("utf-8"),
            "text/html; charset=utf-8",
            "html",
        )
    if fmt in ("md", "markdown"):
        return build_markdown(content, title).encode("utf-8"), "text/markdown; charset=utf-8", "md"
    if fmt == "json":
        return build_json(content).encode("utf-8"), "application/json; charset=utf-8", "json"
    raise ResumeExportError(f"不支持的导出格式: {format!r}，允许 {sorted(VALID_FORMATS)}")
